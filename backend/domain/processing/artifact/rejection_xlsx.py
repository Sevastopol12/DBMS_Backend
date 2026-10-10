"""XLSX rejection artifact generator."""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from backend.domain.processing.models import RowDecision, RowDisposition, SourceDataset
from backend.domain.processing.reserved_columns import (
    COL_DUP_ROLE,
    COL_GROUP_ROWS,
    COL_INVALID_FIELDS,
    COL_ISSUES,
    COL_MERGED_INTO,
    COL_ORIGIN_FILE,
    COL_ORIGIN_ROW,
)

#: Reserved columns appended after the original headers, in fixed order.
ARTIFACT_RESERVED_COLUMNS: tuple[str, ...] = (
    COL_ORIGIN_FILE,
    COL_ORIGIN_ROW,
    COL_ISSUES,
    COL_INVALID_FIELDS,
    COL_DUP_ROLE,
    COL_MERGED_INTO,
    COL_GROUP_ROWS,
)

SHEET_DATA = "data"
SHEET_INSTRUCTIONS = "instructions"

_LIST_SEPARATOR = ";"


def _stringify(value: object | None) -> str:
    if value is None:
        return ""
    return str(value)


def _invalid_fields(issue_codes: Sequence[str]) -> list[str]:
    """Derive field names from ``<KIND>:<field>`` issue codes (best effort)."""

    fields: set[str] = set()
    for code in issue_codes:
        head, sep, tail = code.rpartition(":")
        if sep and head and tail and " " not in tail:
            fields.add(tail)
    return sorted(fields)


def _row_cells(worksheet, values: Sequence[str]) -> None:
    row: list[WriteOnlyCell] = []
    for value in values:
        cell = WriteOnlyCell(worksheet, value=value)
        # Every artifact cell is an explicit string so values like "=1+1"
        # can never be interpreted as formulas by a spreadsheet reader.
        cell.data_type = "s"
        row.append(cell)
    worksheet.append(row)


def build_rejection_xlsx(
    source: SourceDataset,
    decisions: Sequence[RowDecision],
    *,
    file_level_reason: str | None,
    file_id: UUID,
) -> bytes:
    """Build the user-editable rejection workbook from original source values."""

    rows_by_number: dict[int, dict[str, object | None]] = {
        index + 2: row for index, row in enumerate(source.rows)
    }
    rejected = [
        decision
        for decision in decisions
        if decision.disposition is RowDisposition.REJECTED
    ]
    decision_by_row = {decision.source_row_number: decision for decision in rejected}

    if file_level_reason is not None:
        included = sorted(rows_by_number)
    else:
        wanted: set[int] = set(decision_by_row)
        for decision in rejected:
            for member in decision.group_row_numbers:
                if member in rows_by_number:
                    wanted.add(member)
        included = sorted(wanted)

    workbook = Workbook(write_only=True)
    data_sheet = workbook.create_sheet(SHEET_DATA)
    _row_cells(data_sheet, [*source.headers, *ARTIFACT_RESERVED_COLUMNS])

    file_id_text = str(file_id)
    for row_number in included:
        original = rows_by_number[row_number]
        decision = decision_by_row.get(row_number)
        if decision is None and file_level_reason is None:
            for candidate in rejected:
                if row_number in candidate.group_row_numbers:
                    decision = candidate
                    break
        if file_level_reason is not None:
            issue_codes = [file_level_reason]
        elif decision is not None:
            issue_codes = list(decision.issue_codes)
        else:
            issue_codes = []
        group_numbers = (
            sorted(decision.group_row_numbers) if decision is not None else []
        )
        values = [_stringify(original.get(header)) for header in source.headers]
        values.extend(
            [
                file_id_text,
                str(row_number),
                _LIST_SEPARATOR.join(issue_codes),
                _LIST_SEPARATOR.join(_invalid_fields(issue_codes)),
                decision.duplicate_role.value if decision is not None else "NONE",
                (
                    str(decision.merged_into_row)
                    if decision is not None and decision.merged_into_row is not None
                    else ""
                ),
                ",".join(str(number) for number in group_numbers),
            ]
        )
        _row_cells(data_sheet, values)

    instructions = workbook.create_sheet(SHEET_INSTRUCTIONS)
    _row_cells(
        instructions,
        ["How to correct and re-upload this rejection artifact"],
    )
    for line in (
        (
            "Sheet 1 (data) lists every rejected source row with its original "
            "columns, values, and order preserved."
        ),
        (
            "Columns starting with _dbms_ are processing metadata (issue codes, "
            "duplicate relationships, originating row). They are ignored on "
            "re-upload and may be left in place or removed."
        ),
        (
            "Fix the values in the original columns, then re-upload this workbook "
            "through the upload_report/ route. It is ingested as a new file with "
            "a new file id."
        ),
        "Only sheet 1 is read back; this instruction sheet is never parsed.",
    ):
        _row_cells(instructions, [line])

    buffer = BytesIO()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    return buffer.getvalue()


__all__ = [
    "ARTIFACT_RESERVED_COLUMNS",
    "SHEET_DATA",
    "SHEET_INSTRUCTIONS",
    "build_rejection_xlsx",
]
