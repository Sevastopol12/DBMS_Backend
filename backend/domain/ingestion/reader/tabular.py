"""CSV and XLSX implementations of the source dataset reader."""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Sequence
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.domain.models import SourceDataset, SourceRow

from .base import SourceDatasetReader
from .errors import (
    CorruptFileError,
    EmptyFileError,
    EmptyWorksheetError,
    ReaderError,
    UnreadableInputError,
    UnsupportedDelimiterError,
    UnsupportedFormatError,
)


class TabularSourceReader(SourceDatasetReader):
    """Read comma-delimited UTF-8 CSV and first-sheet XLSX files."""

    def read(
        self, file_bytes: bytes, filename: str, source_file_id: UUID
    ) -> SourceDataset:

        if not isinstance(file_bytes, bytes):
            raise UnreadableInputError("Source input must be bytes")

        if not isinstance(filename, str) or not filename.strip():
            raise UnreadableInputError("Source filename is required")

        suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if suffix == "csv":
            return self._read_csv(file_bytes, filename, source_file_id)

        if suffix == "xlsx":
            return self._read_xlsx(file_bytes, filename, source_file_id)

        raise UnsupportedFormatError(f"Unsupported source format for {filename!r}")

    def _read_csv(
        self, file_bytes: bytes, filename: str, source_file_id: UUID
    ) -> SourceDataset:
        if not file_bytes or not file_bytes.strip():
            raise EmptyFileError(f"CSV file {filename!r} is empty")

        try:
            text = file_bytes.decode("utf-8-sig")

        except UnicodeDecodeError as exc:
            raise UnreadableInputError(
                f"CSV file {filename!r} is not valid UTF-8"
            ) from exc

        if not text.strip():
            raise EmptyFileError(f"CSV file {filename!r} is empty")

        try:
            dialect = csv.Sniffer().sniff(text, delimiters=",;\t|")

        except csv.Error:
            dialect = None

        if dialect is not None and dialect.delimiter != ",":
            raise UnsupportedDelimiterError(
                f"CSV file {filename!r} uses unsupported delimiter {dialect.delimiter!r}"
            )

        try:
            records = list(csv.reader(io.StringIO(text, newline=""), strict=True))

        except csv.Error as exc:
            raise CorruptFileError(f"CSV file {filename!r} is malformed") from exc
        if not records or not any(records):
            raise EmptyFileError(f"CSV file {filename!r} is empty")

        headers = self._validate_headers(
            [str(value) for value in records[0]],
            filename=filename,
            empty_error=EmptyFileError,
            source_type="CSV file",
        )
        rows = self._build_rows(records[1:], headers, filename, "CSV file")

        return SourceDataset(
            source_file_id=source_file_id, filename=filename, headers=headers, rows=rows
        )

    def _read_xlsx(
        self, file_bytes: bytes, filename: str, source_file_id: UUID
    ) -> SourceDataset:
        if not file_bytes:
            raise EmptyFileError(f"XLSX file {filename!r} is empty")
        try:
            workbook = load_workbook(
                io.BytesIO(file_bytes), read_only=True, data_only=False
            )
        except (
            InvalidFileException,
            OSError,
            ValueError,
            KeyError,
            zipfile.BadZipFile,
        ) as exc:
            raise CorruptFileError(f"XLSX file {filename!r} is corrupt") from exc

        try:
            worksheet = workbook.worksheets[0]
            records = list(worksheet.iter_rows(values_only=True))
        except (IndexError, OSError, ValueError) as exc:
            raise UnreadableInputError(
                f"XLSX file {filename!r} cannot be read"
            ) from exc
        finally:
            workbook.close()

        if not records:
            raise EmptyWorksheetError(f"First worksheet in {filename!r} is empty")
        headers = self._validate_headers(
            ["" if value is None else str(value) for value in records[0]],
            filename=filename,
            empty_error=EmptyWorksheetError,
            source_type="XLSX file",
        )
        rows = self._build_rows(records[1:], headers, filename, "XLSX file")
        return SourceDataset(
            source_file_id=source_file_id, filename=filename, headers=headers, rows=rows
        )

    @staticmethod
    def _validate_headers(
        headers: list[str],
        *,
        filename: str,
        empty_error: type[ReaderError],
        source_type: str,
    ) -> list[str]:
        """Validate headers while preserving their source spelling and order."""
        if not headers or any(not header.strip() for header in headers):
            raise empty_error(f"{source_type} {filename!r} has no header row")

        if len(set(headers)) != len(headers):
            raise CorruptFileError(
                f"{source_type} {filename!r} has duplicate column names"
            )

        return headers

    @staticmethod
    def _build_rows(
        records: Sequence[Sequence[object]],
        headers: Sequence[str],
        filename: str,
        source_type: str,
    ) -> list[SourceRow]:
        """Map positional records to headers and pad omitted cells with ``None``."""
        rows: list[SourceRow] = []
        for values in records:
            if len(values) > len(headers):
                raise CorruptFileError(
                    f"{source_type} {filename!r} contains values without a source column"
                )
            rows.append(
                {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                }
            )
        return rows


def read_source_dataset(
    file_bytes: bytes, filename: str, source_file_id: UUID
) -> SourceDataset:
    """Convenience entry point for callers that do not need a reader instance."""

    return TabularSourceReader().read(file_bytes, filename, source_file_id)


__all__ = ["TabularSourceReader", "read_source_dataset"]
